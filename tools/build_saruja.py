#!/usr/bin/env python3
"""
SARUJA 21596244 – build dát z Excel histórie (MT4 Pro USD, XAUUSD).

Vstup : saruja-21596244-historia.xlsx (hárok SARUJA: Ticket, Open Time, Type, Volume,
        Symbol, Open Price, Close Time, Close Price, Commission, Swap, Profit USD)
Výstup: saruja-21596244-trades.csv   – čitateľná tabuľka obchodov (;)
        saruja-21596244-data.pine    – Pine polia pre indikátor
        saruja-21596244-data.json    – dáta + štatistiky pre HTML dashboard

Použitie: python3 tools/build_saruja.py saruja-21596244-historia.xlsx .
Časy sú ponechané tak, ako sú v exporte (neprepočítané); časové pásmo si
indikátor kalibruje sám podľa sviečok (alebo ručne cez vstup).
"""
import json
import statistics
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

OZ_PER_LOT = 100.0          # XAUUSD: 1 lot = 100 oz
BASKET_GAP_S = 60           # pozície zatvorené do 60 s od seba = kôš


def parse_dt(s):
    return datetime.strptime(str(s).strip(), "%Y-%m-%d %H:%M:%S")


def ms(dt):
    return int(dt.replace(tzinfo=timezone.utc).timestamp() * 1000)


def fmt_dur(h):
    m = round(h * 60)
    if m < 1:
        return "< 1 min"
    d, rem = divmod(m, 1440)
    hh, mm = divmod(rem, 60)
    if d:
        return f"{d}d {hh}h {mm:02d}m"
    if hh:
        return f"{hh}h {mm:02d}m"
    return f"{mm} min"


def load(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() for h in rows[0]]
    col = {h: i for i, h in enumerate(hdr)}
    trades, balance_ops = [], []
    for r in rows[1:]:
        if r[col["Ticket"]] is None:
            continue
        typ = str(r[col["Type"]]).strip().upper()
        if typ in ("BALANCE", "CREDIT", "DEPOSIT", "WITHDRAWAL"):
            balance_ops.append({
                "ticket": str(r[col["Ticket"]]),
                "time": parse_dt(r[col["Open Time"]]),
                "amount": float(r[col["Profit USD"]]),
            })
            continue
        trades.append({
            "ticket": str(r[col["Ticket"]]),
            "open": parse_dt(r[col["Open Time"]]),
            "type": typ,
            "dir": 1 if typ == "BUY" else -1,
            "lot": float(r[col["Volume"]]),
            "symbol": str(r[col["Symbol"]]),
            "open_price": float(r[col["Open Price"]]),
            "close": parse_dt(r[col["Close Time"]]),
            "close_price": float(r[col["Close Price"]]),
            "commission": float(r[col["Commission"]] or 0),
            "swap": float(r[col["Swap"]] or 0),
            "profit": float(r[col["Profit USD"]]),
        })
    summary = {}
    if len(wb.worksheets) > 1:
        for k, v in wb.worksheets[1].iter_rows(values_only=True):
            if k is not None:
                summary[str(k)] = v
    return trades, balance_ops, summary


def enrich(trades):
    trades.sort(key=lambda t: (t["open"], t["ticket"]))
    for n, t in enumerate(trades, 1):
        t["n"] = n
        t["dur_h"] = (t["close"] - t["open"]).total_seconds() / 3600.0
        t["dur_txt"] = fmt_dur(t["dur_h"])
        t["move"] = round((t["close_price"] - t["open_price"]) * t["dir"], 2)   # USD/oz v prospech obchodu
        t["pips"] = round(t["move"] * 10)
        t["profit_calc"] = round(t["move"] * t["lot"] * OZ_PER_LOT, 2)
        t["profit_per_lot"] = round(t["profit"] / t["lot"], 2) if t["lot"] else None
        t["check_ok"] = abs(t["profit_calc"] + t["commission"] + t["swap"] - t["profit"]) < 0.05 * max(1, t["lot"]) + 0.01
    # koše: zoradené podľa času zatvorenia, medzera ≤ 60 s
    by_close = sorted(trades, key=lambda t: t["close"])
    bid, prev = 0, None
    groups = defaultdict(list)
    for t in by_close:
        if prev is None or (t["close"] - prev).total_seconds() > BASKET_GAP_S:
            bid += 1
        groups[bid].append(t)
        prev = t["close"]
    baskets = []
    k = 0
    for bid in sorted(groups):
        g = groups[bid]
        if len(g) < 2:
            for t in g:
                t["basket"] = 0
            continue
        k += 1
        for t in g:
            t["basket"] = k
        baskets.append({
            "id": k,
            "trades": [t["n"] for t in g],
            "lots": round(sum(t["lot"] for t in g), 2),
            "net": round(sum(t["profit"] for t in g), 2),
            "first_open": min(t["open"] for t in g),
            "close": max(t["close"] for t in g),
            "avg_entry": round(sum(t["open_price"] * t["lot"] for t in g) / sum(t["lot"] for t in g), 2),
        })
    return trades, baskets


def exposure(trades):
    ev = []
    for t in trades:
        ev.append((t["open"], 1, t["lot"]))
        ev.append((t["close"], -1, t["lot"]))
    ev.sort(key=lambda e: (e[0], e[1]))
    run_n = run_l = 0.0
    max_n = max_l = 0.0
    t_max_n = t_max_l = None
    series = []
    for tm, d, lot in ev:
        run_n += d
        run_l += d * lot
        run_l = round(run_l, 2)
        series.append((tm, int(run_n), run_l))
        if run_n > max_n:
            max_n, t_max_n = run_n, tm
        if run_l > max_l:
            max_l, t_max_l = run_l, tm
    return int(max_n), t_max_n, max_l, t_max_l, series


def stats(trades, balance_ops, baskets):
    deposit = sum(b["amount"] for b in balance_ops if b["amount"] > 0)
    withdrawals = sum(b["amount"] for b in balance_ops if b["amount"] < 0)
    profits = [t["profit"] for t in trades]
    wins = [p for p in profits if p > 0]
    losses = [p for p in profits if p < 0]
    gross_p = sum(wins)
    gross_l = -sum(losses)
    net = sum(profits)
    # krivka zostatku podľa zatvorenia (closed-trade DD)
    bal = deposit
    peak = bal
    max_dd = 0.0
    max_dd_pct = 0.0
    curve = [{"t": min(b["time"] for b in balance_ops).strftime("%Y-%m-%d %H:%M:%S") if balance_ops else None, "bal": bal, "n": 0}]
    dd_at = None
    for t in sorted(trades, key=lambda t: (t["close"], t["ticket"])):
        bal = round(bal + t["profit"], 2)
        t["balance_after"] = bal
        peak = max(peak, bal)
        dd = peak - bal
        if dd > max_dd:
            max_dd, max_dd_pct, dd_at = dd, dd / peak * 100 if peak else 0, t["close"]
        curve.append({"t": t["close"].strftime("%Y-%m-%d %H:%M:%S"), "bal": bal, "n": t["n"]})
    durs = [t["dur_h"] for t in trades]
    dur_w = [t["dur_h"] for t in trades if t["profit"] > 0]
    dur_l = [t["dur_h"] for t in trades if t["profit"] < 0]
    max_n, t_max_n, max_l, t_max_l, _ = exposure(trades)
    # série
    best_streak = cur = worst_streak = curl = 0
    for t in sorted(trades, key=lambda t: t["close"]):
        if t["profit"] > 0:
            cur += 1; curl = 0
        elif t["profit"] < 0:
            curl += 1; cur = 0
        best_streak = max(best_streak, cur)
        worst_streak = max(worst_streak, curl)
    days = (max(t["close"] for t in trades) - min(t["open"] for t in trades)).days + 1
    sells = sum(1 for t in trades if t["dir"] < 0)
    return {
        "deposit": deposit,
        "withdrawals": withdrawals,
        "trades": len(trades),
        "wins": len(wins),
        "losses": len(losses),
        "win_rate": round(100 * len(wins) / len(trades), 1),
        "net": round(net, 2),
        "gross_profit": round(gross_p, 2),
        "gross_loss": round(gross_l, 2),
        "profit_factor": round(gross_p / gross_l, 2) if gross_l else None,
        "avg_win": round(statistics.mean(wins), 2) if wins else None,
        "avg_loss": round(statistics.mean(losses), 2) if losses else None,
        "loss_win_ratio": round(-statistics.mean(losses) / statistics.mean(wins), 2) if wins and losses else None,
        "expectancy": round(net / len(trades), 2),
        "largest_win": max(profits),
        "largest_loss": min(profits),
        "roi_pct": round(100 * net / deposit, 1) if deposit else None,
        "final_balance": round(deposit + withdrawals + net, 2),
        "max_closed_dd": round(max_dd, 2),
        "max_closed_dd_pct": round(max_dd_pct, 2),
        "max_closed_dd_at": dd_at.strftime("%Y-%m-%d %H:%M:%S") if dd_at else None,
        "avg_dur_h": round(statistics.mean(durs), 2),
        "med_dur_h": round(statistics.median(durs), 2),
        "max_dur_h": round(max(durs), 2),
        "avg_dur_win_h": round(statistics.mean(dur_w), 2) if dur_w else None,
        "avg_dur_loss_h": round(statistics.mean(dur_l), 2) if dur_l else None,
        "total_lots": round(sum(t["lot"] for t in trades), 2),
        "avg_lot": round(statistics.mean(t["lot"] for t in trades), 3),
        "max_lot": max(t["lot"] for t in trades),
        "max_concurrent": max_n,
        "max_concurrent_at": t_max_n.strftime("%Y-%m-%d %H:%M:%S"),
        "max_concurrent_lots": max_l,
        "max_concurrent_lots_at": t_max_l.strftime("%Y-%m-%d %H:%M:%S"),
        "best_streak": best_streak,
        "worst_streak": worst_streak,
        "baskets": len(baskets),
        "trades_in_baskets": sum(len(b["trades"]) for b in baskets),
        "sell_pct": round(100 * sells / len(trades), 1),
        "days": days,
        "first_open": min(t["open"] for t in trades).strftime("%Y-%m-%d %H:%M:%S"),
        "last_close": max(t["close"] for t in trades).strftime("%Y-%m-%d %H:%M:%S"),
        "balance_curve": curve,
        "all_checks_ok": all(t["check_ok"] for t in trades),
    }


def weekly(trades):
    wk = defaultdict(list)
    for t in trades:
        iso = t["close"].isocalendar()
        monday = (t["close"] - timedelta(days=t["close"].weekday())).date()
        wk[(iso[0], iso[1], monday)].append(t)
    out = []
    for (y, w, monday), g in sorted(wk.items()):
        p = [t["profit"] for t in g]
        out.append({
            "week": f"{y}-W{w:02d}",
            "monday": monday.isoformat(),
            "trades": len(g),
            "wins": sum(1 for x in p if x > 0),
            "net": round(sum(p), 2),
            "lots": round(sum(t["lot"] for t in g), 2),
            "avg_dur_h": round(statistics.mean(t["dur_h"] for t in g), 2),
        })
    return out


def write_csv(trades, path):
    hdr = ["#", "Ticket", "Otvorené", "Smer", "Lot", "Entry", "Zatvorené", "Exit",
           "Pohyb (USD/oz)", "Pip", "Zisk USD", "Zisk USD / lot", "Trvanie (h)", "Trvanie", "Kôš", "Zostatok po"]
    lines = [";".join(hdr)]
    for t in trades:
        lines.append(";".join([
            str(t["n"]), t["ticket"], t["open"].strftime("%d.%m.%Y %H:%M:%S"), t["type"],
            f'{t["lot"]:.2f}', f'{t["open_price"]:.2f}', t["close"].strftime("%d.%m.%Y %H:%M:%S"),
            f'{t["close_price"]:.2f}', f'{t["move"]:.2f}', str(t["pips"]), f'{t["profit"]:.2f}',
            f'{t["profit_per_lot"]:.2f}', f'{t["dur_h"]:.2f}', t["dur_txt"],
            str(t["basket"]) if t["basket"] else "", f'{t["balance_after"]:.2f}',
        ]))
    Path(path).write_text("﻿" + "\n".join(lines) + "\n", encoding="utf-8")


def pine_arr(name, typ, vals):
    return f"var array<{typ}> {name} = array.from({', '.join(vals)})"


def write_pine(trades, baskets, st, path):
    def f2(x):
        s = f"{x:.2f}".rstrip("0").rstrip(".")
        return s if "." in s else s + ".0"
    lines = [
        "//@version=6",
        f"// SARUJA 21596244 – údaje o obchodoch ({len(trades)} obchodov, {st['first_open'][:10]} – {st['last_close'][:10]}, XAUUSD, MT4 Pro USD)",
        "// TK = ticket, TO/TC = čas otvorenia/zatvorenia v ms (čas exportu brané ako UTC – indikátor ho posúva podľa kalibrácie),",
        "// D = smer (1 BUY, -1 SELL), VO = lot, PO/PC = cena otvorenia/zatvorenia, PR = zisk USD, BK = kôš (0 = samostatný)",
        pine_arr("TK", "int", [t["ticket"] for t in trades]),
        pine_arr("TO", "int", [str(ms(t["open"])) for t in trades]),
        pine_arr("TC", "int", [str(ms(t["close"])) for t in trades]),
        pine_arr("D", "int", [str(t["dir"]) for t in trades]),
        pine_arr("VO", "float", [f2(t["lot"]) for t in trades]),
        pine_arr("PO", "float", [f2(t["open_price"]) for t in trades]),
        pine_arr("PC", "float", [f2(t["close_price"]) for t in trades]),
        pine_arr("PR", "float", [f2(t["profit"]) for t in trades]),
        pine_arr("BK", "int", [str(t["basket"]) for t in trades]),
        f"DEPOSIT = {f2(st['deposit'])}",
    ]
    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    xlsx = sys.argv[1]
    out = Path(sys.argv[2] if len(sys.argv) > 2 else ".")
    trades, balance_ops, summary = load(xlsx)
    trades, baskets = enrich(trades)
    st = stats(trades, balance_ops, baskets)
    wk = weekly(trades)
    # verifikácia voči hárku Súhrn
    if summary:
        exp_n = int(summary.get("Tradov", len(trades)))
        exp_p = float(summary.get("Súčet profit USD", st["net"]))
        assert exp_n == len(trades), f"počet obchodov {len(trades)} ≠ súhrn {exp_n}"
        assert abs(exp_p - st["net"]) < 0.01, f"súčet profitu {st['net']} ≠ súhrn {exp_p}"
    assert st["all_checks_ok"], "profit nesedí s (entry-exit)×lot×100 pri niektorom obchode"
    write_csv(trades, out / "saruja-21596244-trades.csv")
    write_pine(trades, baskets, st, out / "saruja-21596244-data.pine")
    _, _, _, _, expo = exposure(trades)
    data = {
        "account": {"number": "21596244", "name": "SARUJA", "platform": str(summary.get("Platforma", "MT4")),
                    "symbol": "XAUUSD", "source": str(summary.get("Zdroj", "")), "time_note": str(summary.get("Časy", ""))},
        "deposits": [{"time": b["time"].strftime("%Y-%m-%d %H:%M:%S"), "amount": b["amount"], "ticket": b["ticket"]} for b in balance_ops],
        "trades": [{k: (v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, datetime) else v) for k, v in t.items()} for t in trades],
        "baskets": [{**b, "first_open": b["first_open"].strftime("%Y-%m-%d %H:%M:%S"), "close": b["close"].strftime("%Y-%m-%d %H:%M:%S")} for b in baskets],
        "exposure": [{"t": tm.strftime("%Y-%m-%d %H:%M:%S"), "n": n, "lots": l} for tm, n, l in expo],
        "weekly": wk,
        "stats": st,
    }
    (out / "saruja-21596244-data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"OK: {len(trades)} obchodov, net {st['net']} USD, {len(baskets)} košov, "
          f"max súčasne {st['max_concurrent']} pozícií / {st['max_concurrent_lots']} lot, "
          f"closed DD {st['max_closed_dd']} USD ({st['max_closed_dd_pct']} %)")


if __name__ == "__main__":
    main()
